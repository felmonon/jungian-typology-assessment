#!/usr/bin/env python3
"""Round-trip and interoperability checks for a Pagewright/WordCraft-compatible CLI.

For every fixture in fixtures/docx it:
  1. validates the input with the Open XML SDK validator (tools/ooxml-validate),
  2. converts DOCX -> DOCX and DOCX -> PDF with the CLI under test,
  3. validates the re-saved DOCX with the Open XML SDK,
  4. compares paragraph text read by python-docx (an independent reader) before and after,
  5. checks structural counts (tables, pictures, hyperlinks, sections, headers/footers, fields),
  6. opens the re-saved DOCX in LibreOffice (an independent consumer) and converts it to PDF,
  7. records PDF page counts from the CLI and from LibreOffice.
Corrupt fixtures must fail cleanly: a non-zero exit with a message, no crash signal, no output.

    python3 tools/roundtrip_check.py --cli path/to/cli --out /tmp/rt [--json report.json]

The CLI must accept `convert IN OUT`. Both `wordcraft-cli` and `pagewright-cli` do.
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
FIX = ROOT / "fixtures" / "docx"
BAD = ROOT / "fixtures" / "corrupt"
VALIDATOR = ROOT / "tools" / "ooxml-validate" / "bin" / "ooxml-validate.dll"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def run(cmd: list[str], timeout: int = 180) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)


def validate(paths: list[Path]) -> dict[str, dict]:
    if not VALIDATOR.exists():
        return {str(p): {"valid": None, "note": "validator not built"} for p in paths}
    res = run(["dotnet", str(VALIDATOR), *map(str, paths)], timeout=300)
    out = {}
    for line in res.stdout.splitlines():
        if line.startswith("{"):
            d = json.loads(line)
            out[d["file"]] = d
    return out


def docx_facts(path: Path) -> dict:
    from docx import Document
    d = Document(str(path))
    paras = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                paras.extend(p.text for p in c.paragraphs)
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        xml = z.read("word/document.xml").decode("utf8")
    return {
        "text": [re.sub(r"\s+", " ", t).strip() for t in paras if t.strip()],
        "tables": len(d.tables),
        "table_rows": sum(len(t.rows) for t in d.tables),
        "pictures": len(d.inline_shapes),
        "hyperlinks": xml.count("<w:hyperlink"),
        "sections": len(d.sections),
        "headers": sum(1 for n in names if re.match(r"word/header\d*\.xml", n)),
        "footers": sum(1 for n in names if re.match(r"word/footer\d*\.xml", n)),
        "fields": len(re.findall(r"(?:instrText[^>]*>\s*(PAGE|NUMPAGES)\b)|(?:fldSimple[^>]*instr=\"\s*(PAGE|NUMPAGES))", " ".join(
            z2 for z2 in _all_story_xml(path)))),
        "comments": "word/comments.xml" in names,
        "ins_del": xml.count("<w:ins ") + xml.count("<w:del "),
        "bookmarks": xml.count("<w:bookmarkStart"),
        "custom_xml": any(n.startswith("customXml/item") and n.endswith(".xml") and "Props" not in n for n in names),
        "repeat_header": xml.count("<w:tblHeader"),
        "parts": sorted(names),
    }


def _all_story_xml(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as z:
        return [z.read(n).decode("utf8", "replace") for n in z.namelist() if re.match(r"word/(document|header\d*|footer\d*)\.xml", n)]


def pdf_pages(path: Path) -> int | None:
    res = run(["pdfinfo", str(path)])
    m = re.search(r"Pages:\s+(\d+)", res.stdout)
    return int(m.group(1)) if m else None


def libreoffice_pdf(docx: Path, outdir: Path) -> Path | None:
    outdir.mkdir(parents=True, exist_ok=True)
    res = run(["soffice", "--headless", "--norestore", "--convert-to", "pdf", "--outdir", str(outdir), str(docx)], timeout=240)
    pdf = outdir / (docx.stem + ".pdf")
    return pdf if res.returncode == 0 and pdf.exists() else None


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--cli", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--json")
    ap.add_argument("--only", nargs="*")
    args = ap.parse_args()
    out = Path(args.out); shutil.rmtree(out, ignore_errors=True); out.mkdir(parents=True)
    fixtures = sorted(FIX.glob("*.docx"))
    if args.only:
        fixtures = [f for f in fixtures if f.stem in args.only]

    report = {"cli": args.cli, "fixtures": {}, "corrupt": {}}
    inputs_valid = validate(fixtures)
    resaved: list[Path] = []
    for f in fixtures:
        r: dict = {"input_valid": inputs_valid.get(str(f), {}).get("valid"),
                   "input_errors": inputs_valid.get(str(f), {}).get("errors", [])[:5]}
        rt = out / f"{f.stem}.roundtrip.docx"
        pdf = out / f"{f.stem}.pdf"
        c1 = run([args.cli, "convert", str(f), str(rt)])
        c2 = run([args.cli, "convert", str(f), str(pdf)])
        r["docx_exit"], r["pdf_exit"] = c1.returncode, c2.returncode
        r["docx_stderr"] = (c1.stderr or "")[-400:]
        r["pdf_stderr"] = (c2.stderr or "")[-400:]
        if c1.returncode == 0 and rt.exists():
            resaved.append(rt)
            a, b = docx_facts(f), docx_facts(rt)
            r["text_equal"] = a["text"] == b["text"]
            if not r["text_equal"]:
                missing = [t for t in a["text"] if t not in b["text"]]
                extra = [t for t in b["text"] if t not in a["text"]]
                r["text_missing"], r["text_extra"] = missing[:5], extra[:5]
            r["structure"] = {k: [a[k], b[k]] for k in a if k not in ("text", "parts")}
            r["structure_changed"] = sorted(k for k, (x, y) in r["structure"].items() if x != y)
            r["parts_dropped"] = sorted(set(a["parts"]) - set(b["parts"]))
            lo = libreoffice_pdf(rt, out / "libreoffice")
            r["libreoffice_opens_resaved"] = lo is not None
            r["libreoffice_pages"] = pdf_pages(lo) if lo else None
            lo_in = libreoffice_pdf(f, out / "libreoffice-input")
            r["libreoffice_pages_input"] = pdf_pages(lo_in) if lo_in else None
        r["cli_pdf_pages"] = pdf_pages(pdf) if pdf.exists() else None
        report["fixtures"][f.stem] = r
    resaved_valid = validate(resaved)
    for rt in resaved:
        stem = rt.name.removesuffix(".roundtrip.docx")
        v = resaved_valid.get(str(rt), {})
        report["fixtures"][stem]["output_valid"] = v.get("valid")
        report["fixtures"][stem]["output_errors"] = v.get("errors", [])[:8]
        report["fixtures"][stem]["output_error_count"] = v.get("errorCount")

    for f in sorted(BAD.glob("*")):
        target = out / f"corrupt-{f.stem}.docx"
        try:
            c = run([args.cli, "convert", str(f), str(target)], timeout=60)
            report["corrupt"][f.name] = {"exit": c.returncode, "crashed": c.returncode < 0 or "panicked" in c.stderr,
                                         "wrote_output": target.exists(), "message": (c.stderr or c.stdout).strip()[-300:]}
        except subprocess.TimeoutExpired:
            report["corrupt"][f.name] = {"exit": None, "crashed": True, "message": "timeout"}

    text = json.dumps(report, indent=2, ensure_ascii=False)
    if args.json:
        Path(args.json).write_text(text)
    # Short human summary
    for name, r in report["fixtures"].items():
        print(f"{name:22} in_valid={r.get('input_valid')!s:5} out_valid={r.get('output_valid')!s:5} "
              f"errs={r.get('output_error_count')!s:4} text_eq={r.get('text_equal')!s:5} "
              f"changed={','.join(r.get('structure_changed', [])) or '-'} dropped={len(r.get('parts_dropped', []))} "
              f"lo_ok={r.get('libreoffice_opens_resaved')!s:5} pages cli/lo_in/lo_out={r.get('cli_pdf_pages')}/"
              f"{r.get('libreoffice_pages_input')}/{r.get('libreoffice_pages')}")
    for name, r in report["corrupt"].items():
        print(f"corrupt {name:28} exit={r['exit']} crashed={r['crashed']} wrote={r.get('wrote_output')} msg={r['message'][:90]!r}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
